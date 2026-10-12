import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
import re

doc = docx.Document()

# Adjust margins
sections = doc.sections
for section in sections:
    section.top_margin = Inches(1.0)
    section.bottom_margin = Inches(1.0)
    section.left_margin = Inches(1.0)
    section.right_margin = Inches(1.0)

# Colors
C_PRIMARY = RGBColor(139, 92, 246) # Violet
C_DARK = RGBColor(22, 30, 49)

with open("c:/quiz/docs/FA2_Project_Report.md", "r", encoding="utf-8") as f:
    lines = f.readlines()

for line in lines:
    line_str = line.strip()
    if not line_str:
        continue
    
    if line_str.startswith("# "):
        h = doc.add_heading(line_str[2:], level=0)
        h.alignment = WD_ALIGN_PARAGRAPH.CENTER
        for run in h.runs:
            run.font.color.rgb = C_PRIMARY
            run.font.size = Pt(22)
            run.font.bold = True
    elif line_str.startswith("## "):
        h = doc.add_heading(line_str[3:], level=1)
        for run in h.runs:
            run.font.color.rgb = C_PRIMARY
            run.font.size = Pt(16)
            run.font.bold = True
    elif line_str.startswith("### "):
        h = doc.add_heading(line_str[4:], level=2)
        for run in h.runs:
            run.font.size = Pt(13)
            run.font.bold = True
    elif line_str.startswith("- "):
        p = doc.add_paragraph(style='List Bullet')
        text = line_str[2:]
        # Simple bold formatting parser
        parts = re.split(r'(\*\*.*?\*\*)', text)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)
    elif line_str.startswith("```"):
        continue
    elif line_str == "---":
        p = doc.add_paragraph()
        r = p.add_run("―" * 45)
        r.font.color.rgb = RGBColor(200, 200, 200)
    else:
        p = doc.add_paragraph()
        parts = re.split(r'(\*\*.*?\*\*)', line_str)
        for part in parts:
            if part.startswith('**') and part.endswith('**'):
                r = p.add_run(part[2:-2])
                r.bold = True
            else:
                p.add_run(part)

doc.save("c:/quiz/docs/FA2_Project_Report.docx")
print("Successfully generated updated FA2_Project_Report.docx!")
