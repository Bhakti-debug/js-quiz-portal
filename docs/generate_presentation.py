import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
import os

prs = pptx.Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

# High-contrast modern dark palette matching the app UI
COLOR_BG = RGBColor(11, 15, 25)
COLOR_CARD = RGBColor(22, 30, 49)
COLOR_CARD_BORDER = RGBColor(139, 92, 246)
COLOR_PRIMARY = RGBColor(167, 139, 250)     # Light Violet
COLOR_CYAN = RGBColor(56, 189, 248)        # Sky Cyan
COLOR_GREEN = RGBColor(52, 211, 153)       # Emerald
COLOR_TEXT_WHITE = RGBColor(248, 250, 252)
COLOR_TEXT_MUTED = RGBColor(148, 163, 184)

def add_header(slide, title_text, subtitle_text=''):
    header = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(0.5), Inches(11.733), Inches(1.2))
    header.fill.solid()
    header.fill.fore_color.rgb = COLOR_CARD
    header.line.color.rgb = COLOR_CARD_BORDER
    header.line.width = Pt(1.5)
    
    tf = header.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title_text
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = COLOR_TEXT_WHITE
    
    if subtitle_text:
        p2 = tf.add_paragraph()
        p2.text = subtitle_text
        p2.font.size = Pt(13)
        p2.font.color.rgb = COLOR_CYAN

def create_base_slide():
    blank_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(blank_layout)
    bg = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
    bg.fill.solid()
    bg.fill.fore_color.rgb = COLOR_BG
    bg.line.color.rgb = COLOR_BG
    return slide

def add_bullets(tf, points):
    for idx, (title, desc) in enumerate(points):
        p = tf.add_paragraph()
        p.text = f"• {title}: "
        p.font.bold = True
        p.font.size = Pt(15)
        p.font.color.rgb = COLOR_CYAN
        run = p.add_run()
        run.text = desc
        run.font.bold = False
        run.font.color.rgb = COLOR_TEXT_WHITE
        if idx < len(points) - 1:
            p_spacer = tf.add_paragraph()
            p_spacer.font.size = Pt(6)

def add_two_column_cards(slide, left_title, left_points, right_title, right_points):
    # Left Card
    card_l = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(5.6), Inches(4.9))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = COLOR_CARD
    card_l.line.color.rgb = COLOR_CARD_BORDER
    tf_l = card_l.text_frame
    tf_l.word_wrap = True
    p_l = tf_l.paragraphs[0]
    p_l.text = left_title
    p_l.font.size = Pt(19)
    p_l.font.bold = True
    p_l.font.color.rgb = COLOR_PRIMARY
    tf_l.add_paragraph().font.size = Pt(6)
    add_bullets(tf_l, left_points)
    
    # Right Card
    card_r = slide.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(6.933), Inches(2.0), Inches(5.6), Inches(4.9))
    card_r.fill.solid()
    card_r.fill.fore_color.rgb = COLOR_CARD
    card_r.line.color.rgb = COLOR_CYAN
    tf_r = card_r.text_frame
    tf_r.word_wrap = True
    p_r = tf_r.paragraphs[0]
    p_r.text = right_title
    p_r.font.size = Pt(19)
    p_r.font.bold = True
    p_r.font.color.rgb = COLOR_GREEN
    tf_r.add_paragraph().font.size = Pt(6)
    add_bullets(tf_r, right_points)

# ==================== SLIDE 1: Title Slide ====================
s1 = create_base_slide()
card1 = s1.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.5), Inches(1.1), Inches(10.333), Inches(5.3))
card1.fill.solid()
card1.fill.fore_color.rgb = COLOR_CARD
card1.line.color.rgb = COLOR_CARD_BORDER
card1.line.width = Pt(2)

tf1 = card1.text_frame
tf1.word_wrap = True

p = tf1.paragraphs[0]
p.text = "JavaScript Quiz & Assessment Portal"
p.font.size = Pt(36)
p.font.bold = True
p.font.color.rgb = COLOR_PRIMARY
p.alignment = PP_ALIGN.CENTER

p = tf1.add_paragraph()
p.text = "Formative Assessment - 2 (FA-2) Project Defense"
p.font.size = Pt(22)
p.font.bold = True
p.font.color.rgb = COLOR_TEXT_WHITE
p.alignment = PP_ALIGN.CENTER

p = tf1.add_paragraph()
p.text = "\nCourse: Advanced Web Technologies (MCA)  |  Academic Year: 2026 - 2027\nClass & Semester: SYMCA, Semester III\nCourse Owner & Evaluator: Dr. Avinash Chormale"
p.font.size = Pt(15)
p.font.color.rgb = COLOR_CYAN
p.alignment = PP_ALIGN.CENTER

p = tf1.add_paragraph()
p.text = "\nCore Stack: React 19 • Vite 8 • React Router v7 • Pure CSS3 Glassmorphism • RBAC Auth"
p.font.size = Pt(13)
p.font.color.rgb = COLOR_TEXT_MUTED
p.alignment = PP_ALIGN.CENTER

# ==================== SLIDE 2: Problem & Objectives ====================
s2 = create_base_slide()
add_header(s2, "1. Problem Statement & Project Objectives", "Engineering a dedicated assessment environment for real-world JavaScript & React developers")
add_two_column_cards(
    s2,
    "The Problem in Developer Assessments",
    [
        ("Shallow Question Banks", "Traditional platforms test simple trivia rather than closures, event loop cycles, and React hooks."),
        ("No Instant Pedagogical Feedback", "Candidates are informed if an answer is wrong but rarely provided with immediate code rationale."),
        ("Lack of Administrative Control", "Standard mock tests do not offer instructors an administrative dashboard to manage question banks.")
    ],
    "Core Project Objectives",
    [
        ("Real-time Timed Assessment Engine", "Execute timed evaluations with automated submission on timer expiry."),
        ("Explainable Feedback System", "Display immediate code explanations for why answers are correct or incorrect."),
        ("Gamified Developer Progression", "Drive engagement with automated XP rewards (10 XP/question) and rank levels."),
        ("Unified Role-Based Portals", "Provide separate workflows for candidate assessments and faculty CRUD operations.")
    ]
)

# ==================== SLIDE 3: FA-1 Foundations Recap ====================
s3 = create_base_slide()
add_header(s3, "2. Phase-I (FA-1) Architectural Foundation", "Key deliverables and frontend building blocks established in Phase-I")
add_two_column_cards(
    s3,
    "FA-1 Milestones Delivered",
    [
        ("Project Scaffolding", "Configured high-performance Vite tooling with zero build overhead and HMR."),
        ("UI/UX Design System", "Designed a custom dark glassmorphism theme using backdrop-filter and CSS variables."),
        ("Component Decomposition", "Built atomic components: navigation bars, quiz summary cards, and stats containers."),
        ("State & Event Listeners", "Employed useState and controlled inputs for dynamic filtering and responsive UI updates.")
    ],
    "FA-1 Technical Outcomes",
    [
        ("Zero CSS Framework Bloat", "Built from scratch without Tailwind or Bootstrap, guaranteeing minimal bundle size."),
        ("Modular File Hierarchy", "Separated concerns across components, pages, context, and mock services."),
        ("Responsive Layouts", "Implemented Flexbox and CSS Grid to support seamless mobile, tablet, and desktop views.")
    ]
)

# ==================== SLIDE 4: FA-2 Advancements & Authentication ====================
s4 = create_base_slide()
add_header(s4, "3. Phase-II (FA-2) Advancements & Authentication", "Elevating the project into a secure, full-capability application with user authentication")
add_two_column_cards(
    s4,
    "Authentication Architecture",
    [
        ("Multi-User Support", "Stores registered users, active sessions, and credential verification via localStorage."),
        ("Sign In & Sign Up Views", "Dedicated /login and /register interfaces with input validation and password confirmation."),
        ("Pre-Seeded Demo Accounts", "Includes one-click demo credentials for Student and Faculty evaluator convenience."),
        ("Session State Synchronization", "Active session persists seamlessly across page reloads and browser restarts.")
    ],
    "Role-Based Access Control (RBAC)",
    [
        ("Candidate / Student Role", "Can take assessments, earn XP, unlock developer ranks, and view attempt history."),
        ("Faculty / Admin Role", "Full access to the Admin Portal (/admin) for CRUD operations and test suite authoring."),
        ("Role Badges & User Status", "Distinct visual indicators in the navbar and dashboard to distinguish roles at a glance.")
    ]
)

# ==================== SLIDE 5: Protected Routes & Security ====================
s5 = create_base_slide()
add_header(s5, "4. Security Architecture & Route Protection", "Defending assessment workflows and administrative workbenches from unauthorized access")
box5 = s5.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.9))
box5.fill.solid()
box5.fill.fore_color.rgb = COLOR_CARD
box5.line.color.rgb = COLOR_CARD_BORDER
tf5 = box5.text_frame
tf5.word_wrap = True
p5 = tf5.paragraphs[0]
p5.text = "ProtectedRoute Component Implementation"
p5.font.size = Pt(19)
p5.font.bold = True
p5.font.color.rgb = COLOR_PRIMARY
tf5.add_paragraph().font.size = Pt(6)

points5 = [
    ("Authentication Verification", "Intercepts route requests and verifies active session. Unauthenticated users are redirected to /login with original location state preserved for post-login return."),
    ("Granular Role Verification (requireAdmin)", "Guards /admin routes. Non-admin users are blocked and presented with a clear 'Administrator Access Required' access-denied view."),
    ("SPA Client-Side Route Safety", "Integrated vercel.json and _redirects rules to ensure bookmarking and direct URL refreshes work seamlessly without 404 errors."),
    ("Session Termination & Logout", "Single-click logout cleans session state immediately and redirects users safely to the authentication screen.")
]
add_bullets(tf5, points5)

# ==================== SLIDE 6: System Design & State Flow ====================
s6 = create_base_slide()
add_header(s6, "5. System Architecture & State Management", "Decoupled SPA design utilizing React Context API and simulated async persistence")
add_two_column_cards(
    s6,
    "Global Context (QuizContext.jsx)",
    [
        ("Single Source of Truth", "Centralizes authentication state, user stats, question banks, and assessment history."),
        ("No Prop Drilling", "Provides custom hook useQuiz() accessible across all pages and deep components."),
        ("Dynamic Re-rendering", "Updates XP, Level progression, and score records instantly upon assessment completion.")
    ],
    "Simulated Service Layer (api.js)",
    [
        ("Async/Await Promises", "Simulates realistic REST API delays (250-300ms) to model real-world client-server architectures."),
        ("LocalStorage Persistence", "Saves quizzes, users, and attempt records reliably across browser sessions."),
        ("Plug-and-Play Swapping", "Can easily be swapped with real Axios or Fetch endpoints without modifying UI code.")
    ]
)

# ==================== SLIDE 7: Timed Assessment Engine ====================
s7 = create_base_slide()
add_header(s7, "6. Timed Assessment Engine & Gamification", "Custom timer hook, instant code pedagogy, and automated XP leveling system")
add_two_column_cards(
    s7,
    "Custom Timer Hook (useTimer.js)",
    [
        ("Accurate Interval Countdown", "Encapsulates useEffect and useRef to guarantee consistent countdown execution."),
        ("Automatic Submission", "Triggers automated completion callback when timeLeft === 0, preventing expired attempts."),
        ("Formatted Display", "Converts raw seconds into clean MM:SS readouts with color-coded warning states.")
    ],
    "Gamification Dynamics",
    [
        ("XP Rewards", "Awards 10 XP per correct answer upon test submission."),
        ("Level Formula", "Automated advancement: Level = floor(XP / 100) + 1."),
        ("Instant Pedagogical Feedback", "Highlights correct/incorrect selections with code explanations immediately on confirmation.")
    ]
)

# ==================== SLIDE 8: Admin CRUD Portal & Validation ====================
s8 = create_base_slide()
add_header(s8, "7. Admin CRUD Portal & Form Validation", "Comprehensive question bank management with rigorous client-side input validations")
add_two_column_cards(
    s8,
    "Admin CRUD Features (/admin)",
    [
        ("Create Quiz", "Initiate new assessment sets with custom category, difficulty, and time limits."),
        ("Read Questions", "Visualize question trees, answer keys, and options in a clean interactive view."),
        ("Update Question", "Modifies question text, option choices, and explanations in-place."),
        ("Delete Operations", "Remove obsolete questions or entire quiz collections with confirmation safeguards.")
    ],
    "Form Field Validation Rules",
    [
        ("Non-Empty Fields", "Enforces non-blank question text and explanations."),
        ("Option Boundary Check", "Requires at least 2 non-empty option strings."),
        ("Correct Option Guard", "Validates that the selected correct answer index aligns with valid option bounds."),
        ("Sanitized Submissions", "Trims whitespace to protect dataset integrity.")
    ]
)

# ==================== SLIDE 9: Testing & Build Optimization ====================
s9 = create_base_slide()
add_header(s9, "8. Verification, Testing & Build Metrics", "Production-grade bundling, automated linting, and build optimization results")
box9 = s9.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.9))
box9.fill.solid()
box9.fill.fore_color.rgb = COLOR_CARD
box9.line.color.rgb = COLOR_CARD_BORDER
tf9 = box9.text_frame
tf9.word_wrap = True
p9 = tf9.paragraphs[0]
p9.text = "Production Build & Quality Assurance Results"
p9.font.size = Pt(19)
p9.font.bold = True
p9.font.color.rgb = COLOR_PRIMARY
tf9.add_paragraph().font.size = Pt(6)

points9 = [
    ("Build Execution (Vite Rollup)", "Compiled 1,836 modules into optimized chunks in ~500ms with zero errors or warnings."),
    ("Asset Minification", "Generated lightweight production bundles: index.html (0.45 kB), CSS (9.67 kB), and JS chunks (gzip ~87 kB)."),
    ("Static Code Analysis (Oxlint)", "Scanned all source code modules with 104 rules; clean passes on unused variables and lifecycle checks."),
    ("Deployment Readiness", "Equipped with SPA fallback routing configurations ready for single-click deployment on Vercel or Netlify.")
]
add_bullets(tf9, points9)

# ==================== SLIDE 10: Conclusion & Project Links ====================
s10 = create_base_slide()
add_header(s10, "9. Conclusion, Deployment Links & Defense Q&A", "Ready for university evaluation and viva examination")
box10 = s10.shapes.add_shape(pptx.enum.shapes.MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(2.0), Inches(11.733), Inches(4.9))
box10.fill.solid()
box10.fill.fore_color.rgb = COLOR_CARD
box10.line.color.rgb = COLOR_CYAN
tf10 = box10.text_frame
tf10.word_wrap = True
p10 = tf10.paragraphs[0]
p10.text = "Project Links & Evaluation Deliverables"
p10.font.size = Pt(19)
p10.font.bold = True
p10.font.color.rgb = COLOR_GREEN
tf10.add_paragraph().font.size = Pt(6)

points10 = [
    ("GitHub Repository", "[To be added: https://github.com/your-username/quiz]"),
    ("Live Deployed Application", "[To be added: https://your-quiz-app.vercel.app]"),
    ("One-Click Student Demo", "student@quiz.com / student123 (Experience candidate test taking, timer, and XP progression)"),
    ("One-Click Admin Demo", "admin@quiz.com / admin123 (Experience faculty quiz administration & CRUD modifications)"),
    ("Future Enhancements", "Monaco Code Editor integration for live code challenges, WebSocket-based multiplayer rooms.")
]
add_bullets(tf10, points10)

out_path = "c:/quiz/docs/FA2_Project_Presentation.pptx"
prs.save(out_path)
print(f"Presentation successfully saved to: {out_path}")
